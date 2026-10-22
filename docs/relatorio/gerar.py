"""Gera relatorio.docx e relatorio.pdf (python-docx + LibreOffice).

    python3 docs/relatorio/gerar.py

Duas passadas: a primeira gera o PDF para descobrir em que página cai cada título,
figura e tabela; a segunda preenche sumário e listas com esses números.
Requer: python-docx, soffice (LibreOffice), pdftotext (poppler), rsvg-convert.
"""
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import diagramas

AQUI = Path(__file__).resolve().parent
FIG = AQUI / "fig"
RAIZ = AQUI.parent.parent
LARGURA_TEXTO = 16.0  # cm (A4 com margens 3/2)

# ---------------------------------------------------------------------------
# Conteúdo

AUTORES = [
    "EDUARDO DE SOUZA MONTEIRO 212123",
    "FELLIPE ALESSANDRO SCRUPH 222378",
    "GUILHERME MATHEUS FONSECA WAIZBART 211899",
    "JOÃO VICTOR DE OLIVEIRA 212106",
    "RENAN VIZOTO FERREIRA 222220",
    "VINICIUS CÂNDIDO DAS CHAGAS 211873",
]
TITULO = "MODELAMENTO DO CONTROLE EM MALHA FECHADA DA MISTURA AR-COMBUSTÍVEL AUTOMOTIVA"

RESUMO = (
    "Este trabalho apresenta o desenvolvimento de uma bancada simulada para o controle em malha "
    "fechada da mistura ar-combustível de um motor de combustão interna, a partir da leitura de uma "
    "sonda lambda. A unidade de controle eletrônico (ECU) é um firmware em C, com HAL e FreeRTOS, "
    "para o microcontrolador STM32F072, executado sem hardware físico no emulador Renode. O motor, o "
    "atraso de transporte dos gases no escapamento e a dinâmica da sonda são modelados por uma planta "
    "escrita em Go, na qual cada sensor é uma rotina concorrente, sincronizada com o emulador a cada "
    "passo de simulação. A ECU lê os sensores por conversor analógico-digital com DMA, calcula o tempo "
    "de injeção por uma tabela de alimentação com correção proporcional-integral e gera o pulso do "
    "injetor por temporizador; o estado do sistema é publicado em barramento CAN e exibido em um "
    "painel gráfico nativo. Nos ensaios, a tarefa de controle cumpriu o período de 10 ms com jitter "
    "máximo de 10 µs, tempo de execução de cerca de 160 µs e nenhum prazo perdido, e a mistura "
    "retornou a λ = 1,00 ± 0,02 em 1,3 s, sem sobressinal, após uma redução de 20 % na vazão do "
    "injetor. A bancada permite estudar, de forma segura e reprodutível, o controle em malha fechada "
    "sob restrições de tempo real."
)
PALAVRAS = "Sistemas de tempo real. FreeRTOS. Sonda lambda. Controle PI. STM32."

ABSTRACT = (
    "This work presents a simulated test bench for closed-loop control of the air-fuel mixture of an "
    "internal combustion engine based on a lambda sensor reading. The engine control unit (ECU) is C "
    "firmware, built on the vendor HAL and FreeRTOS, for the STM32F072 microcontroller, running "
    "without physical hardware on the Renode emulator. The engine, the exhaust gas transport delay "
    "and the sensor dynamics are modeled by a plant written in Go, in which each sensor is a "
    "concurrent routine synchronized with the emulator at every simulation step. The ECU samples the "
    "sensors through an analog-to-digital converter with DMA, computes the injection time from a "
    "fuel map with proportional-integral correction and drives the injector pulse with a hardware "
    "timer; the system state is published on a CAN bus and shown on a native graphical panel. In the "
    "tests, the control task met its 10 ms period with a maximum jitter of 10 µs, an execution time "
    "of about 160 µs and no missed deadlines, and the mixture returned to λ = 1.00 ± 0.02 within "
    "1.3 s, without overshoot, after a 20 % reduction in injector flow. The bench provides a safe and "
    "reproducible platform for studying closed-loop control under real-time constraints."
)
KEYWORDS = "Real-time systems. FreeRTOS. Lambda sensor. PI control. STM32."


def texto(doc):
    """Elementos textuais, na ordem. Cada item é (tipo, dados...)."""
    c = []
    a = c.append

    # 1 -------------------------------------------------------------------
    a(("h1", "1 INTRODUÇÃO"))
    a(("h2", "1.1 INTRODUÇÃO DO GRUPO"))
    a(("p", "O presente trabalho foi desenvolvido pelos alunos Eduardo de Souza Monteiro, Fellipe "
            "Alessandro Scruph, Guilherme Matheus Fonseca Waizbart, João Victor de Oliveira, Renan "
            "Vizoto Ferreira e Vinicius Cândido das Chagas, estudantes do curso de Engenharia de "
            "Computação da Facens. O projeto faz parte das atividades práticas da disciplina de "
            "Tópicos Especiais de Engenharia, orientada pelos professores Adilson Rocha e Marcos "
            "Fábio Jardini."))
    a(("h2", "1.2 DESCRIÇÃO DA ÁREA DO TEMA, RELEVÂNCIA E PONTOS CRÍTICOS"))
    a(("p", "O tema escolhido aborda sistemas embarcados aplicados à área automotiva, com foco no "
            "controle eletrônico da injeção de combustível. Manter a proporção correta entre ar e "
            "combustível, indicada pelo fator lambda medido pela sonda, é essencial para que o veículo "
            "polua menos, economize combustível e funcione de maneira estável. Na prática, esse "
            "controle enfrenta desafios importantes, como a diferença de resposta entre sensores "
            "comuns e sensores de precisão e o tempo que os gases da queima levam para sair dos "
            "cilindros e alcançar o sensor no escapamento. Esse atraso varia com a rotação do motor e "
            "pode desestabilizar o controle se o sistema não for rápido e bem ajustado."))
    a(("h3", "1.2.1 Descrição dos objetivos específicos e escopo"))
    a(("p", "O objetivo do grupo é construir uma bancada de testes simulada para monitorar a sonda "
            "lambda e ajustar a mistura em malha fechada sem um motor real. O firmware da ECU é "
            "escrito para o STM32 e executado no emulador Renode; o modelo do motor, o atraso do "
            "escapamento (buffer circular) e a resposta da sonda (equação de primeira ordem) ficam em "
            "uma planta simulada externa, que se comunica com o microcontrolador emulado pelas "
            "entradas analógicas e pela saída do injetor. Os cálculos de tempo concentram-se na taxa "
            "de amostragem do conversor analógico, no período e no jitter da tarefa de controle e na "
            "largura do pulso de injeção em microssegundos. Com esses dados, analisa-se quanto tempo o "
            "sistema leva para corrigir a mistura após uma perturbação e se o controle permanece "
            "estável."))
    a(("h2", "1.3 CRONOGRAMA GERAL DE DESENVOLVIMENTO (MACRO TAREFAS)"))
    a(("tabela", "Cronograma de desenvolvimento", ["Etapa", "Macro Tarefa", "Descrição / Entregável"], [
        ["1", "Levantamento Teórico e Equações",
         "Definição das fórmulas da mistura, do atraso do escapamento e da resposta da sonda lambda."],
        ["2", "Configuração dos Periféricos do STM32",
         "Configuração, no emulador Renode, dos timers de disparo, da leitura analógica por DMA, do "
         "pulso do injetor e do CAN."],
        ["3", "Programação da Planta Simulada",
         "Implementação, em Go, do modelo do motor, do atraso do escapamento e da sonda, "
         "sincronizado com o microcontrolador emulado."],
        ["4", "Implementação do Código de Controle",
         "Desenvolvimento da lógica de malha fechada para calcular a correção do tempo de injeção."],
        ["5", "Testes de Resposta e Coleta de Tempos",
         "Medição do tempo de resposta da malha fechada, avaliação de atrasos e validação do "
         "sistema na bancada."],
        ["6", "Escrita e Entrega do Relatório",
         "Organização dos gráficos gerados, análise das medições de tempo e elaboração do "
         "documento final."],
    ], [1.6, 4.4, 10.0], False))

    # 2 -------------------------------------------------------------------
    a(("h1", "2 RESUMO DE SISTEMAS OPERACIONAIS"))
    a(("p", "Um sistema operacional (SO) é a camada de software situada entre o hardware e as "
            "aplicações. Ele cumpre dois papéis: oferecer abstrações mais simples que o hardware "
            "(arquivos, processos, dispositivos) e gerenciar os recursos do computador, como "
            "processador, memória, tempo e dispositivos de entrada e saída, entre programas que "
            "competem por eles (TANENBAUM; BOS, 2016)."))
    a(("p", "A unidade de execução gerenciada pelo SO é o processo, isto é, um programa em execução "
            "com seu próprio espaço de endereçamento, registradores e estado. Um processo pode conter "
            "várias threads, que compartilham a memória e são escalonadas de forma independente. O "
            "escalonador decide qual tarefa ocupa o processador, alternando-as entre os estados "
            "pronto, executando e bloqueado. Como tarefas concorrentes acessam dados comuns, o SO "
            "fornece mecanismos de sincronização e comunicação, como semáforos, mutexes e filas de "
            "mensagens, que protegem as seções críticas e evitam condições de corrida."))
    a(("p", "O acesso aos dispositivos é feito por drivers, módulos que encapsulam os registradores e "
            "as interrupções de cada periférico e expõem uma interface uniforme ao restante do sistema. "
            "Em microcontroladores, esse papel costuma ser desempenhado por bibliotecas do fabricante, "
            "como a HAL (Hardware Abstraction Layer) da STMicroelectronics usada neste projeto."))
    a(("p", "Um sistema operacional de tempo real (RTOS) é aquele em que a correção depende não só do "
            "resultado, mas também do instante em que ele é produzido (KOPETZ, 2011). Os prazos podem "
            "ser rígidos (hard), quando perdê-los causa falha; firmes (firm), quando o resultado "
            "atrasado perde o valor; ou flexíveis (soft), quando apenas degradam o desempenho. Por "
            "isso, um RTOS privilegia o determinismo em vez da vazão média: escalonamento preemptivo "
            "por prioridade fixa, latência de interrupção limitada, serviços de tempo de execução "
            "previsível e pequeno consumo de memória. A análise clássica de escalonabilidade de Liu e "
            "Layland (1973) mostra que n tarefas periódicas com prioridades atribuídas pela taxa "
            "(menor período, maior prioridade) cumprem seus prazos se a utilização total do "
            "processador não exceder n(2^{1/n} − 1)."))
    a(("h2", "2.1 RESUMO DE IMPLEMENTAÇÃO DE UM RTOS: FREERTOS"))
    a(("p", "O FreeRTOS é um núcleo de tempo real de código aberto (licença MIT), portado para mais de "
            "quarenta arquiteturas e com poucos kilobytes de código (BARRY, 2016). A aplicação é "
            "dividida em tarefas, cada uma com pilha e prioridade próprias, nos estados executando, "
            "pronta, bloqueada ou suspensa. O escalonador é preemptivo: a tarefa pronta de maior "
            "prioridade sempre executa, e tarefas de mesma prioridade revezam a cada tick. O tick é "
            "uma interrupção periódica que serve de base para atrasos e temporizações, como a função "
            "vTaskDelayUntil, que mantém um período fixo sem acumular erro."))
    a(("p", "Para comunicação e sincronização, o FreeRTOS oferece filas (cópia de dados por valor, "
            "seguras entre tarefas e interrupções), semáforos binários e contadores, mutexes com "
            "herança de prioridade e notificações diretas de tarefa, uma alternativa mais leve ao "
            "semáforo que não exige objeto adicional. Rotinas de interrupção usam as variantes "
            "FromISR das funções e solicitam troca de contexto com portYIELD_FROM_ISR. Trechos "
            "curtos que não podem ser interrompidos são protegidos por taskENTER_CRITICAL e "
            "taskEXIT_CRITICAL. Neste projeto, o núcleo roda no Cortex-M0 a 48 MHz, com preempção "
            "habilitada, tick de 1 kHz, cinco níveis de prioridade e alocador heap_4 com 7 KB."))

    # 3 -------------------------------------------------------------------
    a(("h1", "3 DESCRIÇÃO DO PROCESSO MODELADO"))
    a(("p", "A qualidade da mistura é expressa pelo fator lambda, razão entre a relação ar-combustível "
            "real e a estequiométrica (AFR_{st} = 14,7 para a gasolina pura):"))
    a(("eq", "λ = m_{ar} / (AFR_{st} · m_{comb})"))
    a(("p", "Com λ = 1 a queima é completa e o catalisador de três vias opera com máxima eficiência; "
            "λ < 1 indica mistura rica e λ > 1, pobre (HEYWOOD, 2018; BOSCH, 2005). O processo "
            "modelado é a cadeia injetor → combustão → escapamento → sonda → ECU. O motor simulado "
            "tem quatro cilindros e 1,6 L. A massa de ar por ciclo segue o modelo alfa-N (borboleta e "
            "rotação), e a de combustível depende do tempo de abertura do bico t_{i}, descontado o "
            "tempo morto de abertura:"))
    a(("eq", "m_{ar} = m_{ar,max} · f(TPS) · η(RPM) + m_{vaz}"))
    a(("eq", "m_{comb} = k_{inj} · Q_{inj} · (t_{i} − t_{morto})"))
    a(("p", "Nessas equações, Q_{inj} = 3 mg/ms, t_{morto} = 0,8 ms e m_{ar,max} = 480 mg. Os termos "
            "m_{vaz} (vazamento de ar não medido) e k_{inj} (bico sujo ou desgastado) são as "
            "perturbações aplicadas nos ensaios. O gás leva um tempo T_{d} para chegar à sonda, que "
            "responde como um sistema de primeira ordem com τ = 100 ms:"))
    a(("eq", "T_{d} = 180 / RPM + 0,020 s        τ · dλ_{s}/dt = λ(t − T_{d}) − λ_{s}"))
    a(("p", "A ECU estima a parcela de combustível por uma tabela de alimentação t_{base}(TPS, RPM) "
            "de 7 × 8 pontos, com interpolação bilinear, e corrige o erro residual com um controlador "
            "proporcional-integral (OGATA, 2010) sobre o erro e = λ_{med} − λ_{alvo}:"))
    a(("eq", "t_{i} = t_{morto} + t_{base} · (1 + c),     c = K_{p} · e + K_{i} · ∫ e dt"))
    a(("p", "com K_{p} = 0,15, K_{i} = 2,0 s^{−1} e |c| ≤ 0,30. Um anti-windup condicional impede "
            "que o integrador cresça enquanto a saída estiver saturada. Em malha aberta, c = 0."))
    a(("p", "A Figura 1 mostra a arquitetura da bancada. O Renode emula o STM32F072 executando o mesmo "
            "binário que seria gravado na placa. A planta em Go é o mestre do tempo: a cada 5 ms de "
            "tempo simulado, escreve as tensões dos sensores nos canais do ADC, avança a emulação e lê "
            "a largura do pulso programada no temporizador do injetor, como faria uma ponta de "
            "osciloscópio. O controlador CAN emulado é ligado à interface virtual vcan0 do Linux, por "
            "onde o painel recebe o estado da ECU e da planta e envia comandos de acelerador, "
            "perturbação e modo de malha."))
    a(("fig", "Arquitetura da bancada simulada", FIG / "arquitetura.png", 15.5))
    a(("p", "As Figuras 2 e 3 detalham os dois laços. Na ECU, a amostragem é feita por hardware "
            "(temporizador, ADC e DMA) e só o fim de cada bloco de 10 amostras acorda a tarefa de "
            "controle. Na planta, a física avança em passos de 1 ms e a sincronização com a ECU "
            "ocorre a cada cinco passos."))
    a(("fig2", ("Fluxograma da tarefa de controle (ECU)", FIG / "fluxo-controle.png"),
              ("Fluxograma da planta simulada", FIG / "fluxo-planta.png")))

    a(("h2", "3.1 CÁLCULOS DE TEMPOS E GARANTIAS DE ATUAÇÃO"))
    a(("p", "A Tabela 2 reúne as janelas de tempo de cada etapa. A malha de lambda é lenta, dominada "
            "pelo atraso do escapamento (50 a 245 ms) e pela constante da sonda (100 ms), de modo que "
            "um período de controle de 10 ms amostra a planta de 5 a 25 vezes mais rápido que seu "
            "atraso. O prazo de cada ativação da tarefa de controle é o próprio período: o cálculo "
            "deve terminar antes que o DMA complete o bloco seguinte."))
    a(("tabela", "Janelas de tempo e metas por etapa",
       ["Etapa", "Período / janela", "Meta ou garantia", "Classe"], [
        ["Amostragem (TIM15 → ADC → DMA)", "1 ms por conjunto de 3 canais",
         "Disparo por hardware, sem uso da CPU", "Hard (hardware)"],
        ["ControlTask: média, PI e TIM2", "10 ms (bloco de 10 amostras)",
         "Jitter ≤ 100 µs; execução ≤ 1 ms; nenhum prazo perdido", "Firm"],
        ["Pulso do injetor (TIM2, 1 µs)", "1 ciclo do motor: 20 a 150 ms",
         "Resolução de 1 µs; t_{i} ≤ 90 % do ciclo", "Hard (hardware)"],
        ["CanTask: estado, tempos e comandos", "20 ms (tempos a cada 100 ms)",
         "Sem prazo rígido", "Soft"],
        ["Malha de lambda", "Atraso de 50 a 245 ms",
         "λ = 1,00 ± 0,02 em ≤ 2 s; sobressinal ≤ 5 %", "Soft"],
    ], [4.6, 3.8, 5.0, 2.6], True))
    a(("p", "Utilização da CPU. Com execução medida de C ≈ 0,16 ms para T = 10 ms, a tarefa de "
            "controle ocupa U = C/T ≈ 1,6 % do processador. Mesmo supondo que a CanTask consuma 1 ms "
            "a cada 20 ms (5 %), a utilização total (≈ 6,6 %) fica muito abaixo do limite de Liu e "
            "Layland para duas tarefas, 2(√2 − 1) ≈ 82,8 %. As prioridades seguem a mesma regra: a "
            "tarefa de menor período (controle, 10 ms) tem a maior prioridade."))
    a(("p", "Pulso de injeção. O motor de quatro tempos injeta uma vez a cada duas voltas, logo o "
            "ciclo vale T_{ciclo} = 120/RPM: 60 ms a 2000 rpm e 20 ms a 6000 rpm. A 2000 rpm e 23 % "
            "de borboleta, t_{i} ≈ 3,5 ms (5,8 % do ciclo). No pior caso, a 6000 rpm, com borboleta "
            "toda aberta e correção de +30 %, t_{i} ≈ 12,2 ms, ou 61 % do ciclo; o limite de 90 % é "
            "apenas uma proteção. O TIM2 conta a 1 MHz, então a resolução de 1 µs corresponde a "
            "menos de 0,04 % do combustível injetado, desprezível diante da faixa de ±2 % em λ."))
    a(("p", "Latência e barramento. O valor aplicado ao injetor é a média das amostras dos últimos "
            "10 ms e é atualizado no TIM2 em até jitter + execução (≈ 0,17 ms) após o fim do bloco. "
            "A 500 kbit/s, um quadro CAN de 8 bytes ocupa cerca de 0,25 ms; com três a quatro "
            "quadros a cada 20 ms, a carga do barramento fica abaixo de 5 %."))
    a(("p", "Medição. A ECU registra, com um temporizador de 1 µs (TIM7), o instante de liberação de "
            "cada bloco e o início e o fim de cada ativação. Desses instantes saem o período, o "
            "jitter (desvio do período nominal), o tempo de execução e os prazos perdidos (ativações "
            "acumuladas ou término após o próximo bloco), publicados no quadro CAN 0x110. Os tempos "
            "são do relógio virtual do Renode, que não reproduz ciclo a ciclo o processador real. A "
            "Tabela 3 compara as metas com os valores medidos."))
    a(("tabela", "Metas de tempo real e valores medidos",
       ["Métrica", "Meta", "Medido", "Ensaio"], [
        ["Período da ControlTask", "10 ms", "10 000 µs", "Renode e ponta a ponta"],
        ["Jitter da ControlTask", "≤ 100 µs", "1 a 10 µs", "Renode e ponta a ponta"],
        ["Execução da ControlTask", "≤ 1 ms", "158 a 161 µs", "Renode e ponta a ponta"],
        ["Prazos perdidos", "0", "0", "Renode e ponta a ponta"],
        ["Resolução do pulso", "1 µs", "1 µs", "Renode"],
        ["Acomodação (bico −20 %, 2000 rpm)", "≤ 2 s", "1,30 s", "Ponta a ponta"],
        ["Sobressinal", "≤ 5 %", "0 %", "Ponta a ponta"],
    ], [5.4, 2.6, 3.0, 5.0], True))

    # 4 -------------------------------------------------------------------
    a(("h1", "4 DESENVOLVIMENTO"))
    a(("h2", "4.1 TAREFAS E RECURSOS DO FREERTOS"))
    a(("p", "O firmware tem duas tarefas e duas rotinas de interrupção relevantes (Tabela 4). A "
            "ControlTask é acordada por notificação de tarefa, usada como um semáforo contador leve: "
            "a interrupção do DMA incrementa o contador e a tarefa o consome com ulTaskNotifyTake. Se "
            "o valor lido for maior que 1, houve ativações acumuladas, que são contadas como prazos "
            "perdidos. A recepção CAN usa uma fila de oito quadros: a interrupção apenas copia o "
            "quadro para a fila, e a CanTask a esvazia em seu ciclo de 20 ms. Os dados trocados entre "
            "as tarefas ficam em uma estrutura compartilhada protegida por seção crítica. Como a "
            "cópia leva poucos microssegundos, mascarar as interrupções nesse trecho é mais simples e "
            "previsível que usar um mutex, que acrescentaria trocas de contexto."))
    a(("tabela", "Tarefas, interrupções e mecanismos de sincronização",
       ["Elemento", "Prioridade", "Ativação", "Recurso do FreeRTOS", "Função"], [
        ["ISR do DMA", "Hardware", "Fim de bloco (10 ms)", "vTaskNotifyGiveFromISR",
         "Acordar a ControlTask"],
        ["ControlTask", "3", "Notificação", "ulTaskNotifyTake; seção crítica",
         "Média, tabela + PI, TIM2, medição de tempos"],
        ["ISR de recepção CAN", "Hardware", "Quadro recebido", "xQueueSendFromISR (fila de 8)",
         "Entregar comandos à CanTask"],
        ["CanTask", "2", "Periódica, 20 ms", "xTaskDelayUntil; xQueueReceive; seção crítica",
         "Publicar estado e tempos; tratar comandos; log serial"],
        ["Idle", "0", "Tempo livre", "-", "Tarefa ociosa do núcleo"],
    ], [2.8, 2.3, 2.6, 4.5, 3.8], True))
    a(("h2", "4.2 TRECHOS DO CÓDIGO-FONTE"))
    a(("p", "O Código 1 mostra a passagem da interrupção do DMA para a tarefa de controle. O período "
            "da tarefa é imposto pelo hardware, e não por um atraso de software, o que elimina a "
            "deriva entre a amostragem e o cálculo."))
    a(("code", "Interrupção do DMA e laço da ControlTask (adc_dma.c e app.c)",
       """/* adc_dma.c: fim de um bloco de 10 conjuntos de amostras */
BaseType_t woken = pdFALSE;
vTaskNotifyGiveFromISR(notify_task, &woken);
portYIELD_FROM_ISR(woken);

/* app.c: ControlTask, prioridade 3 */
for(;;) {
    uint32_t n = ulTaskNotifyTake(pdTRUE, portMAX_DELAY);
    control_step(n > 1u ? n - 1u : 0u);   /* n > 1: ativação perdida */
}"""))
    a(("p", "O Código 2 mostra a fila de recepção CAN e o laço periódico da CanTask, cuja cadência "
            "fixa de 20 ms é garantida por xTaskDelayUntil."))
    a(("code", "Fila de recepção CAN e laço da CanTask (canbus.c e app.c)",
       """/* canbus.c: interrupção de recepção do bxCAN */
frame.id = header.StdId;
frame.len = (uint8_t)header.DLC;
xQueueSendFromISR(rx_queue, &frame, &woken);

/* app.c: CanTask, prioridade 2 */
for(;;) {
    while(canbus_receive(&rx, 0)) handle_command(&rx);  /* esvazia a fila */
    taskENTER_CRITICAL(); s = shared; taskEXIT_CRITICAL();
    /* ... envia 0x100 e 0x101; a cada 100 ms, 0x110 (tempos) ... */
    xTaskDelayUntil(&wake, pdMS_TO_TICKS(CAN_PERIOD_MS));
}"""))
    a(("p", "O Código 3 traz o núcleo do controlador PI com anti-windup condicional: o integrador só "
            "avança enquanto a saída não ultrapassa o limite de ±30 %."))
    a(("code", "Passo do controlador PI (controller.c)",
       """static float pi_step(controller_t *c, float erro)
{
    float p = PI_KP * erro;
    float i = c->integral + PI_KI * erro * CONTROL_DT_S;
    if(erro > 0.0f && p + i > PI_LIMITE)
        i = maxf(c->integral, PI_LIMITE - p);
    else if(erro < 0.0f && p + i < -PI_LIMITE)
        i = minf(c->integral, -PI_LIMITE - p);
    c->integral = clampf(i, -PI_LIMITE, PI_LIMITE);
    return clampf(p + c->integral, -PI_LIMITE, PI_LIMITE);
}"""))
    a(("p", "Na planta, cada sensor é uma goroutine (thread leve do Go). A cada passo, o estado do "
            "motor é enviado por canal a todos os sensores, e uma barreira (sync.WaitGroup) espera "
            "que todos respondam antes de prosseguir (Código 4), de modo que as leituras ficam "
            "coerentes no tempo simulado."))
    a(("code", "Barreira de sincronização dos sensores (sensors/bank.go)",
       """func (b *Bank) Step(dt float64, st engine.State) []float64 {
    b.wg.Add(len(b.sensors))
    for _, c := range b.in {
        c <- tick{dt: dt, st: st}   // acorda cada goroutine de sensor
    }
    b.wg.Wait()                     // barreira: todos leram o passo
    return append([]float64(nil), b.out...)
}"""))
    a(("h2", "4.3 INTERFACE, TESTES E PRIMEIROS RESULTADOS"))
    a(("p", "O painel (Figura 4), escrito em Go com a biblioteca gráfica raylib, ocupa uma única tela "
            "e lê o barramento CAN. Ele mostra uma barra de mistura rica/ideal/pobre com o valor de λ, "
            "a rotação, o tempo de injeção, a correção do PI, o modo da malha e um gráfico de λ nos "
            "últimos 10 s. Pelo teclado, o visitante acelera o motor, aplica as perturbações de bico "
            "e de vazamento e alterna entre malha aberta e fechada."))
    a(("fig", "Painel do operador", RAIZ / "docs" / "painel.png", 13.0))
    a(("p", "A verificação foi feita em três níveis, todos executados automaticamente a cada envio ao "
            "repositório. Nos testes unitários, o controle em C é testado com o framework Unity, e a "
            "planta, com testes Go. Na integração, o mesmo código C do controle roda dentro da planta "
            "Go (software-in-the-loop), e o firmware completo roda no Renode com casos em Robot "
            "Framework. No teste ponta a ponta, ECU, planta e testes conversam pelo vcan0. A Figura 5 "
            "mostra a resposta a uma redução de 20 % na vazão do bico a 2000 rpm: a mistura empobrece "
            "após o atraso do escapamento, o PI eleva o tempo de injeção em cerca de 25 % e λ volta à "
            "faixa de ±0,02 em 1,3 s, sem sobressinal."))
    a(("fig", "Resposta a uma redução de 20 % na vazão do bico (2000 rpm, malha fechada)",
       FIG / "degrau.png", 13.5))

    # 5 -------------------------------------------------------------------
    a(("h1", "5 CONSIDERAÇÕES FINAIS"))
    a(("p", "A bancada atingiu as metas propostas: a tarefa de controle cumpriu o período de 10 ms "
            "com folga de mais de 98 % do processador e jitter dez vezes menor que o limite, e a "
            "malha corrigiu as perturbações em menos de 2 s. A separação entre amostragem por "
            "hardware, tarefa de controle acordada por notificação e comunicação em tarefa de menor "
            "prioridade manteve o caminho crítico curto e previsível. Como limitação, os tempos foram "
            "medidos no relógio virtual do emulador; a próxima etapa natural é gravar o mesmo binário "
            "em uma placa STM32 e repetir as medições com instrumentos reais."))

    a(("ref", [
        "BARRY, R. *Mastering the FreeRTOS real time kernel*: a hands-on tutorial guide. "
        "[S. l.]: Real Time Engineers Ltd., 2016.",
        "BOSCH. *Manual de tecnologia automotiva*. 25. ed. São Paulo: Blucher, 2005.",
        "HEYWOOD, J. B. *Internal combustion engine fundamentals*. 2. ed. New York: McGraw-Hill, 2018.",
        "KOPETZ, H. *Real-time systems*: design principles for distributed embedded applications. "
        "2. ed. New York: Springer, 2011.",
        "LIU, C. L.; LAYLAND, J. W. Scheduling algorithms for multiprogramming in a hard-real-time "
        "environment. *Journal of the ACM*, v. 20, n. 1, p. 46-61, 1973.",
        "OGATA, K. *Engenharia de controle moderno*. 5. ed. São Paulo: Pearson, 2010.",
        "RENODE. *Renode documentation*. Antmicro. Disponível em: https://renode.readthedocs.io. "
        "Acesso em: 22 out. 2026.",
        "STMICROELECTRONICS. *RM0091*: STM32F0x1/STM32F0x2/STM32F0x8 reference manual. Genebra: "
        "STMicroelectronics, 2017.",
        "TANENBAUM, A. S.; BOS, H. *Sistemas operacionais modernos*. 4. ed. São Paulo: Pearson, 2016.",
    ]))
    return c


# ---------------------------------------------------------------------------
# Formatação

def run_rico(par, s, tam=None, fonte=None):
    """Adiciona texto com _{sub}, ^{sup} e *itálico*."""
    for tok in re.split(r"(_\{[^}]*\}|\^\{[^}]*\}|\*[^*]+\*)", s):
        if not tok:
            continue
        if tok.startswith("_{"):
            r = par.add_run(tok[2:-1]); r.font.subscript = True
        elif tok.startswith("^{"):
            r = par.add_run(tok[2:-1]); r.font.superscript = True
        elif tok.startswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True
        else:
            r = par.add_run(tok)
        if tam:
            r.font.size = Pt(tam)
        if fonte:
            r.font.name = fonte
    return par


def espacamento(par, antes=0, depois=6, linha=1.5):
    pf = par.paragraph_format
    pf.space_before = Pt(antes)
    pf.space_after = Pt(depois)
    pf.line_spacing = linha


def sombrear(par, cor="F2F2F2"):
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), cor)
    ppr.append(shd)


def sombrear_celula(cel, cor="E7E6E6"):
    tcpr = cel._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), cor)
    tcpr.append(shd)


def tabela_inteira(t, larguras=None):
    """Larguras reais das colunas (LibreOffice lê o tblGrid) e linhas que não se dividem."""
    if larguras:
        t.autofit = False
        for col, w in zip(t.columns, larguras):
            col.width = Cm(w)
    linhas = t.rows
    for i, row in enumerate(linhas):
        trpr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); trpr.append(cs)
        if i < len(linhas) - 1:
            for cel in row.cells:
                for p in cel.paragraphs:
                    p.paragraph_format.keep_with_next = True


def campo(par, instr):
    r = par.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1); r._r.append(it); r._r.append(f2)


def margens(sec):
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.left_margin = Cm(3), Cm(3)
    sec.bottom_margin, sec.right_margin = Cm(2), Cm(2)
    sec.header_distance = Cm(1.2)


def cabecalho_logo(sec):
    sec.header.is_linked_to_previous = False
    p = sec.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG / "logo-facens.png"), width=Cm(7.5))


class Relatorio:
    def __init__(self, paginas=None):
        self.pg = paginas or {}
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name = "Arial"
        st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        self.nfig = self.ntab = self.ncod = 0
        self.figs, self.tabs, self.titulos = [], [], []
        margens(self.doc.sections[0])
        cabecalho_logo(self.doc.sections[0])

    # pré-textuais ---------------------------------------------------------
    def centro(self, txt, negrito=True, tam=12, antes=0, depois=12):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        espacamento(p, antes, depois, 1.5)
        r = p.add_run(txt); r.bold = negrito; r.font.size = Pt(tam)
        return p

    def quebra(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def capa(self):
        self.centro("SOROCABA, SP", depois=0)
        self.centro("2026", depois=60)
        for a in AUTORES:
            self.centro(a, depois=6)
        self.centro(TITULO, tam=14, antes=90, depois=100)
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(8)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        espacamento(p, 0, 0, 1.5)
        p.add_run("Documento de Projeto RTOS para a disciplina TEECII.\n"
                  "Orientadores: Adilson Rocha e Marcos Fábio Jardini.")

    def resumo(self, titulo, txt, rotulo, chaves):
        self.quebra()
        self.centro(titulo, depois=18)
        self.par(txt)
        p = self.doc.add_paragraph()
        espacamento(p, 12, 0)
        r = p.add_run(rotulo + ": "); r.bold = True
        p.add_run(chaves)

    def lista(self, titulo, itens):
        self.quebra()
        self.centro(titulo, depois=18)
        for txt, nivel, negrito in itens:
            p = self.doc.add_paragraph()
            espacamento(p, 0, 2, 1.5)
            p.paragraph_format.left_indent = Cm(0.5 * nivel)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(LARGURA_TEXTO), WD_TAB_ALIGNMENT.RIGHT,
                                                      WD_TAB_LEADER.DOTS)
            r = p.add_run(f"{txt}\t{self.pg.get(txt, '')}")
            r.bold = negrito
            r.font.size = Pt(11)

    # textuais -------------------------------------------------------------
    def par(self, txt):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        espacamento(p, 0, 8, 1.5)
        run_rico(p, txt)
        return p

    def titulo(self, txt, nivel, nova_pagina=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.page_break_before = nova_pagina
        if nivel == 1:
            espacamento(p, 18, 12, 1.5)
            r = p.add_run(txt); r.bold = True; r.font.size = Pt(14)
        else:
            espacamento(p, 12, 6, 1.5)
            p.add_run(txt)
        self.titulos.append((txt, nivel))

    def legenda(self, txt):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        espacamento(p, 6, 4, 1.0)
        run_rico(p, txt, tam=10)

    def fonte(self):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        espacamento(p, 2, 10, 1.0)
        r = p.add_run("Fonte: Autoria própria (2026)."); r.font.size = Pt(10)

    def figura(self, txt, caminho, largura):
        self.nfig += 1
        nome = f"Figura {self.nfig} - {txt}"
        self.figs.append(nome)
        self.legenda(nome)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        espacamento(p, 0, 0, 1.0)
        p.add_run().add_picture(str(caminho), width=Cm(largura))
        self.fonte()

    def figuras_lado(self, *figs):
        t = self.doc.add_table(rows=1, cols=len(figs))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cel, (txt, caminho) in zip(t.rows[0].cells, figs):
            self.nfig += 1
            nome = f"Figura {self.nfig} - {txt}"
            self.figs.append(nome)
            p = cel.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            espacamento(p, 0, 4, 1.0)
            run_rico(p, nome, tam=10)
            p2 = cel.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.add_run().add_picture(str(caminho), width=Cm(7.4))
            p3 = cel.add_paragraph()
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p3.add_run("Fonte: Autoria própria (2026)."); r.font.size = Pt(10)
        tabela_inteira(t, [LARGURA_TEXTO / len(figs)] * len(figs))

    def tabela(self, txt, cab, linhas, larguras, com_fonte):
        self.ntab += 1
        nome = f"Tabela {self.ntab} - {txt}"
        self.tabs.append(nome)
        self.legenda(nome)
        t = self.doc.add_table(rows=1 + len(linhas), cols=len(cab))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, linha in enumerate([cab] + linhas):
            for j, val in enumerate(linha):
                cel = t.cell(i, j)
                cel.width = Cm(larguras[j])
                p = cel.paragraphs[0]
                espacamento(p, 2, 2, 1.0)
                run_rico(p, val, tam=10)
                if i == 0:
                    for r in p.runs:
                        r.bold = True
                    sombrear_celula(cel)
        tabela_inteira(t, larguras)
        if com_fonte:
            self.fonte()
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def equacao(self, txt):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        espacamento(p, 2, 8, 1.2)
        run_rico(p, txt, fonte="Cambria Math")

    def codigo(self, txt, cod):
        self.ncod += 1
        self.legenda(f"Código {self.ncod} - {txt}")
        p = self.doc.add_paragraph()
        espacamento(p, 0, 0, 1.0)
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.keep_together = True
        p.paragraph_format.keep_with_next = True
        sombrear(p)
        linhas = cod.split("\n")
        for i, l in enumerate(linhas):
            r = p.add_run(l)
            r.font.name = "Liberation Mono"
            r.font.size = Pt(8.5)
            if i < len(linhas) - 1:
                r.add_break()
        self.fonte()

    def referencias(self, refs):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.page_break_before = True
        espacamento(p, 0, 18, 1.5)
        r = p.add_run("REFERÊNCIAS"); r.bold = True; r.font.size = Pt(14)
        self.titulos.append(("REFERÊNCIAS", 1))
        for ref in refs:
            q = self.doc.add_paragraph()
            espacamento(q, 0, 12, 1.0)
            run_rico(q, ref)

    # montagem -------------------------------------------------------------
    def montar(self):
        corpo = texto(self.doc)
        # Pré-calcula títulos, figuras e tabelas para as listas.
        tit = [(e[1], {"h1": 0, "h2": 1, "h3": 2}[e[0]]) for e in corpo if e[0] in ("h1", "h2", "h3")]
        tit.append(("REFERÊNCIAS", 0))
        figs, tabs, nf, nt = [], [], 0, 0
        for e in corpo:
            if e[0] == "fig":
                nf += 1; figs.append(f"Figura {nf} - {e[1]}")
            elif e[0] == "fig2":
                for f in e[1:]:
                    nf += 1; figs.append(f"Figura {nf} - {f[0]}")
            elif e[0] == "tabela":
                nt += 1; tabs.append(f"Tabela {nt} - {e[1]}")

        self.capa()
        self.resumo("RESUMO", RESUMO, "Palavras-chave", PALAVRAS)
        self.resumo("ABSTRACT", ABSTRACT, "Keywords", KEYWORDS)
        self.lista("LISTA DE FIGURAS", [(f, 0, False) for f in figs])
        self.lista("LISTA DE TABELAS", [(t, 0, False) for t in tabs])
        self.lista("SUMÁRIO", [(t, n, n == 0) for t, n in tit])

        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        margens(sec)
        sec.header.is_linked_to_previous = False
        hp = sec.header.paragraphs[0]
        for r in list(hp.runs):
            r._r.getparent().remove(r._r)
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        campo(hp, "PAGE")
        pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1")
        sec._sectPr.append(pg)

        for e in corpo:
            k = e[0]
            if k == "h1":
                self.titulo(e[1], 1)
            elif k in ("h2", "h3"):
                self.titulo(e[1], 2)
            elif k == "p":
                self.par(e[1])
            elif k == "eq":
                self.equacao(e[1])
            elif k == "fig":
                self.figura(*e[1:])
            elif k == "fig2":
                self.figuras_lado(*e[1:])
            elif k == "tabela":
                self.tabela(*e[1:])
            elif k == "code":
                self.codigo(*e[1:])
            elif k == "ref":
                self.referencias(e[1])
        return figs, tabs, tit


def paginas(pdf):
    """Mapeia título/figura/tabela → número de página impresso."""
    with tempfile.TemporaryDirectory() as d:
        txt = Path(d) / "r.txt"
        subprocess.run(["pdftotext", "-layout", str(pdf), str(txt)], check=True)
        pags = txt.read_text(encoding="utf-8").split("\f")[:-1]  # o último \f fecha a última página
    inicio = next(i for i, p in enumerate(pags)
                  if i > 0 and re.search(r"^\s*1 INTRODUÇÃO\s*$", p, re.M) and "SUMÁRIO" not in p)
    norm = lambda s: re.sub(r"\s+", " ", s)
    return [norm(p) for p in pags], inicio


def gerar():
    diagramas.gerar(FIG)
    docx = AQUI / "relatorio.docx"
    pg = {}
    for passada in range(2):
        rel = Relatorio(pg)
        figs, tabs, tit = rel.montar()
        rel.doc.save(docx)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(AQUI),
                        str(docx)], check=True, capture_output=True)
        pags, inicio = paginas(AQUI / "relatorio.pdf")
        pg = {}
        for nome in [t for t, _ in tit] + figs + tabs:
            m = re.match(r"(Figura|Tabela) \d+ -", nome)
            alvo = m.group(0) if m else re.sub(r"\s+", " ", nome)
            for i in range(inicio, len(pags)):
                if alvo in pags[i]:
                    pg[nome] = i - inicio + 1
                    break
        faltando = [n for n in [t for t, _ in tit] + figs + tabs if n not in pg]
        if faltando:
            print("sem página:", faltando)
    print(f"{docx.name} e relatorio.pdf: {len(pags)} páginas ({len(pags) - inicio} textuais)")


if __name__ == "__main__":
    gerar()
